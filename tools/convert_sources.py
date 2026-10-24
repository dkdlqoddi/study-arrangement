#!/usr/bin/env python3
"""
원본 강의 자료(machine_learning_for_health_science/, git 제외) → 공개용 Markdown(docs/) 변환 파이프라인.

    python3 tools/convert_sources.py                 # 전체 변환 + 인덱스 생성
    python3 tools/convert_sources.py --only week03   # 특정 주차만 변환
    python3 tools/convert_sources.py --index-only    # 변환 생략, manifest.json 기반으로 README/INDEX만 재생성
    python3 tools/convert_sources.py --no-ocr        # 이미지 기반 슬라이드 OCR 생략(빠른 확인용)

변환 규칙 요약
  * PDF   → 글꼴 크기로 제목 단계 복원, 표는 markdown 표로, 이미지 전용 페이지는 OCR(rapidocr)
  * PPTX  → 슬라이드별 제목/본문/표/발표자 노트
  * DOCX  → 문단/제목/표
  * HTML  → markdownify
  * TXT(받아쓰기) → 5분 단위 소제목을 붙인 markdown
  * 모든 결과물은 MAX_TOKENS(o200k_base 기준) 이하가 되도록 제목 경계에서 분할
"""
from __future__ import annotations

import argparse
import datetime as dt
import glob
import html
import json
import os
import re
import shutil
import sys
import warnings
from collections import Counter
from pathlib import Path

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "machine_learning_for_health_science"
OUT = ROOT / "docs"
MANIFEST = OUT / "manifest.json"
TODAY = dt.date.today().isoformat()

MAX_TOKENS = 12000          # ChatGPT 웹에서 한 번에 읽기 편한 크기(o200k_base 토큰 기준)
HEADER_RESERVE = 400        # write_doc 이 각 파트 앞에 붙이는 머리말 몫
GITHUB_REPO = "dkdlqoddi/study-arrangement"

# 공개 범위(2026-10-24 결정): 시험대비 학습자료 요약본과 모아보기(compiled)만 docs/ 에 두어 GitHub에 올린다.
# 그 외(전체판·강의노트·슬라이드·받아쓰기·논문과제·코드·공통 자료)는 private/docs/ 에 두고 .gitignore 로 제외한다.
PUBLIC_CATEGORIES = {"study_guide_summary", "compiled"}
PRIV = ROOT / "private"
PRIV_OUT = PRIV / "docs"
RAW_BASE = f"https://raw.githubusercontent.com/{GITHUB_REPO}/main/"

# ---------------------------------------------------------------------------
# 주차 정의 (원본 폴더 구조가 불규칙해서 손으로 매핑)
#   slides  : (원본 파일명, 출력 slug, 비고)  — 같은 파일이 여러 주차에 있으면 처음 등장한 주차에서만 변환
#   reuse   : 다른 주차에서 이미 변환된 자료(링크만 표시)
#   skip    : 변환하지 않는 원본과 사유
# ---------------------------------------------------------------------------
WEEKS = [
    dict(
        n=1, folder="20260902", date="2026-09-02", slug="week01_20260902",
        title="과목 소개 · Chap1 Introduction · Chap2 Part I (Statistical Learning)",
        topics=["통계학 vs 기계학습(Two cultures)", "supervised/unsupervised/reinforcement", "f 추정, Prediction vs Inference",
                "손실함수와 최적 예측자 E[Y|X]", "Parametric vs Non-parametric", "차원의 저주", "유연성-해석력 trade-off",
                "수업 운영(중간·기말, 팀 프로젝트, AI 도구 사용 보고)"],
        summary_pdf="1차시 시험대비 학습자료.pdf", full_pdf="1차시 시험대비 학습자료(전체판).pdf",
        notes=[("1차시 강의노트.pdf", None)],
        slides=[("Intro (2).pdf", "intro_course_overview", "과목 소개 슬라이드(이미지 위주 → OCR)"),
                ("Chap1 (1).pdf", "chap1_introduction", "ISLR Ch.1 (이미지 위주 → OCR)"),
                ("Chap2_Part1 (1).pdf", "chap2_part1_statistical_learning", "ISLR Ch.2 Part I"),
                ("Chap2_Part2 (1).pdf", "chap2_part2_assessing_model_accuracy", "ISLR Ch.2 Part II")],
        code=[],
        skip=[("Int Statistical Rev - 2021 - Efron - Prediction  Estimation  and Attribution.pdf",
               "학술지 논문 원문(저작권) — docs/common/references.md 의 DOI 링크 참고")],
    ),
    dict(
        n=2, folder="20260909", date="2026-09-09", slug="week02_20260909",
        title="Chap2 Part I 복습(f 추정, parametric/non-parametric, KNN·차원의 저주) · Chap2 Part II(MSE, Bias-Variance, Bayes classifier) · Efron 논문 발표",
        topics=["f를 왜·어떻게 추정하나", "parametric vs non-parametric", "KNN과 차원의 저주", "예측 정확도 vs 해석력", "MSE, training vs test error", "Bias-Variance trade-off", "0-1 loss, Bayes classifier, KNN 분류", "Efron(2020) Prediction·Estimation·Attribution 발표", "Shmueli(2010) To Explain or to Predict? 동료 발표"],
        summary_pdf=None, full_pdf="2차시 시험대비 학습자료(전체판).pdf",
        notes=[("20260909.pdf", None)],
        slides=[("Chap03_Lec.pdf", "chap3_linear_regression", "ISLR Ch.3"),
                ("Ch04-Classification.pdf", "chap4_classification", "ISLR Ch.4 (이미지 위주 → OCR)"),
                ("To Explain or to Predict*발표자료.pdf", "student_presentation_shmueli_to_explain_or_to_predict",
                 "동료 학생의 Shmueli(2010) 논문 발표자료 — 공개 전 동의 확인 필요")],
        code=[("kNN.txt", "knn_regression_demo.R", "KNN 회귀 vs 선형회귀 R 데모")],
        paper_assignment=True,
        skip=[("Int Statistical Rev - 2021 - Efron - Prediction  Estimation  and Attribution.pdf", "논문 원문(저작권)"),
              ("20260909_논문_발표자료_*.pdf", "20260909_발표PPT_v2.1.pptx 의 PDF 출력본과 동일 내용 — pptx 쪽을 변환"),
              ("논문읽기과제/20260909_발표PPT_v1.1.pptx", "발표 초안(최종본 v2.1 변환)"),
              ("논문읽기과제/20260909_발표PPT_v1.2.pptx", "발표 초안"),
              ("논문읽기과제/20260909_발표PPT_v2.1.pptx", "최종본과 다른 중간 버전"),
              ("논문읽기과제/20260909_발표PPT_v2.1 - 복사본.pptx", "중복 복사본"),
              ("논문읽기과제/발제슬라이드_Efron_2020.pptx", "발표 초안(AI 생성)"),
              ("논문읽기과제/발표슬라이드_Efron_2020_5분.pptx", "발표 초안(AI 생성)"),
              ("논문읽기과제/template.pptx", "빈 템플릿"),
              ("논문읽기과제/Efron_2020_Prediction_Estimation_Attribution_완전요약본.pdf", "같은 내용의 _source.html 을 변환"),
              ("논문읽기과제/Efron_2020_Prediction_Estimation_Attribution_요약본_3쪽.pdf", "같은 내용의 _source.html 을 변환"),
              ("논문읽기과제/Efron_예측_추정_귀속_논문_해설_요약.pdf", "같은 내용의 .docx 를 변환"),
              ("논문읽기과제/발제문_Efron_2020.pdf", "같은 내용의 _source.html 을 변환")],
    ),
    dict(
        n=3, folder="20260916", date="2026-09-16", slug="week03_20260916",
        title="Chap2 Part II 마무리 · Double descent 발표 · Chap3 Linear Regression(두 관점·likelihood·추론·변수선택·KNN 비교) · 회귀→경사하강→신경망 · Chap4 분류 시작(로지스틱, MLE·cross-entropy)",
        topics=["모델 정확도 평가 복습", "Double descent(Belkin 2019)", "선형회귀의 두 관점", "likelihood", "추론·변수선택·R²·KNN 비교", "제곱오차 → 경사하강 → 신경망(One Idea, Many Models)", "분류 문제의 정의와 risk", "선형회귀 vs 로지스틱 회귀(link)", "MLE → cross-entropy, 0-1 loss와 surrogate", "다변량 정규분포 보충 덱"],
        summary_pdf="3차시 시험대비 학습자료.pdf", full_pdf="3차시 시험대비 학습자료(전체판).pdf",
        notes=[("20260916.pdf", None)],
        slides=[("Multivariate_Gaussian_Distribution.pdf", "multivariate_gaussian_distribution", "다변량 정규분포 보충"),
                ("regression_to_neural_nets.pdf", "regression_to_neural_nets", "One Idea, Many Models: 회귀 → 신경망"),
                ("logistic_to_neural_nets.pdf", "logistic_to_neural_nets", "로지스틱 회귀 → 신경망")],
        code=[],
        reuse=[("Chap2_Part2.pdf", "week01_20260902/slides/chap2_part2_assessing_model_accuracy.md")],
        skip=[],
    ),
    dict(
        n=4, folder="20260923", date="2026-09-23", slug="week04_20260923",
        title="로지스틱 회귀 심화(우도·경사법·Newton–Raphson·odds ratio·confounding·환자-대조군) · 다항 로지스틱 · 판별분석의 출발점 · TRIPOD/PROBAST 발표",
        topics=["로지스틱 회귀 모형과 링크 함수", "0-1 손실 → 우도 → −log 우도 → 경사법", "표준오차·z·p-value, Newton–Raphson", "odds ratio, confounding", "유병률 vs 환자 비율, 코호트 vs 환자-대조군", "환자-대조군 로지스틱 회귀(선택지표 S)", "다항 로지스틱(baseline category logit)", "판별분석의 출발점(Bayes 정리)", "complete separation", "TRIPOD·PROBAST 논문 발표"],
        summary_pdf="4차시 시험대비 학습자료.pdf", full_pdf="4차시 시험대비 학습자료(전체판).pdf",
        notes=[("20260923.pdf", None)],
        slides=[("Binary_Classifier_Performance.pdf", "binary_classifier_performance", "이진 분류기 성능 지표 (5~6차시 진행)")],
        code=[("NaiveBayes.txt", "naive_bayes_default.R", "e1071::naiveBayes, ISLR Default 데이터 R 예제")],
        reuse=[("Multivariate_Gaussian_Distribution.pdf", "week03_20260916/slides/multivariate_gaussian_distribution.md"),
               ("logistic_to_neural_nets.pdf", "week03_20260916/slides/logistic_to_neural_nets.md"),
               ("regression_to_neural_nets.pdf", "week03_20260916/slides/regression_to_neural_nets.md")],
        skip=[],
    ),
    dict(
        n=5, folder="20260930", date="2026-09-30", slug="week05_20260930",
        title="판별분석(LDA, p=1 → p>1) · 다변량 정규분포 · Naive Bayes·Laplace 평활 · 이진 분류기 성능(정확도 함정·임계값·민감도·특이도·ROC·AUC)",
        topics=["판별분석: p(Y|X) ⇕ p(X|Y)", "LDA p=1 판별점수·모수 추정", "p>1 다변량 정규, GEE·mixed effect 언급", "Multivariate Gaussian(밀도식, Σ, 기하)", "naive Bayes, Laplace 평활", "K=3 분류, Fisher 붓꽃", "정확도의 함정, 오류 비용, 임계값", "민감도·특이도, ROC, AUC"],
        summary_pdf="5차시 시험대비 학습자료.pdf", full_pdf="5차시 시험대비 학습자료(전체판).pdf",
        notes=[("20260930.pdf", None)],
        slides=[("CrossValidation.pdf", "cross_validation", "ISLR Ch.5 CV (5차시에 배포, 6~7차시에 진행 · 일부 페이지 OCR)"),
                ("bootstrap_en.pdf", "bootstrap", "Bootstrap 보충 자료 (6차시에 진행)")],
        code=[],
        reuse=[("Binary_Classifier_Performance.pdf", "week04_20260923/slides/binary_classifier_performance.md"),
               ("Multivariate_Gaussian_Distribution.pdf", "week03_20260916/slides/multivariate_gaussian_distribution.md"),
               ("logistic_to_neural_nets.pdf", "week03_20260916/slides/logistic_to_neural_nets.md"),
               ("regression_to_neural_nets.pdf", "week03_20260916/slides/regression_to_neural_nets.md"),
               ("NaiveBayes.txt", "week04_20260923/code/naive_bayes_default.R")],
        skip=[("음성 260930_183439.m4a / 음성 260930_194512.m4a / 음성 260930_204737.m4a", "Quick Share 폴더의 녹음과 동일(중복)")],
    ),
    dict(
        n=6, folder="20261007", date="2026-10-07", slug="week06_20261007",
        title="분류기 성능 심화(c-통계량=AUC 증명·discrimination vs calibration·PPV/NPV) · Naive Bayes 재정리 · Bootstrap(원리·CI·회귀·OOB) · Cross-Validation 시작(검증셋·k-fold·LOOCV)",
        topics=["ROC·AUC 복습", "c-통계량, AUC = c-통계량 증명", "discrimination vs calibration", "AUC의 한계", "민감도·특이도 vs PPV·NPV(유병률)", "학습 손실 vs 평가지표", "naive Bayes 재정리", "부트스트랩 원리·알고리즘·신뢰구간", "회귀에서의 부트스트랩, (1−1/n)ⁿ, OOB", "교차검증 대원칙, 검증셋, k-fold, LOOCV"],
        summary_pdf="6차시 시험대비 학습자료.pdf", full_pdf="6차시 시험대비 학습자료(전체판).pdf",
        notes=[("20261007.pdf", None)],
        slides=[("ModelSelection_Regularization.pdf", "model_selection_regularization", "ISLR Ch.6 (6차시 배포, 7~8차시 진행 · 이미지 위주 → OCR)"),
                ("AIC_BIC_Derivations.pdf", "aic_bic_derivations", "AIC·BIC 유도 보충 (8차시 진행)"),
                ("OtherPenaltyFunctions.pdf", "other_penalty_functions", "기타 벌점 함수 (8차시 진행 · 이미지 위주 → OCR)")],
        code=[],
        reuse=[],
        skip=[],
    ),
    dict(
        n=7, folder="20261014", date="2026-10-14", slug="week07_20261014",
        title="Cross-Validation 완결(k-fold·LOOCV 지름길·1-SE 규칙·올바른 CV) · Model Selection(subset selection·Cp/AIC/BIC/adjR²·AIC vs BIC 철학) · 축소법 시작(능형회귀·라쏘)",
        topics=["CV의 목적·검증셋 접근", "k-fold CV 식", "LOOCV 지름길 공식(hᵢ), k의 편향–분산", "CV 곡선, 분류 문제의 CV, 1-SE 규칙", "올바른 CV vs 틀린 CV(p=5000, n=50)", "모형 선택 세 방법군, AIC vs BIC 철학", "best subset·forward·backward", "Cp·AIC·BIC·adjusted R²", "자코비안과 변환한 y의 AIC", "축소의 발상, 능형회귀(λ의 CV 선택)", "라쏘(ℓ1, 변수 선택)", "PCR·PLS(교재, 미진행)"],
        summary_pdf="7차시 시험대비 학습자료.pdf", full_pdf="7차시 시험대비 학습자료(전체판).pdf",
        notes=[("20261014.pdf", None)],
        slides=[],
        code=[],
        reuse=[("CrossValidation.pdf", "week05_20260930/slides/cross_validation.md"),
               ("ModelSelection_Regularization.pdf", "week06_20261007/slides/model_selection_regularization.md")],
        skip=[("20261014-notes.pdf", "다른 과목(Topics in Biochemical Engineering, 2026-10-13) 노트가 잘못 들어온 파일 — 변환 제외")],
    ),
    dict(
        n=8, folder="20261022", date="2026-10-21", slug="week08_20261022",
        title="AIC·BIC 유도(KL 발산·Laplace 근사) · 능형회귀·라쏘 심화(해 경로·표준화·sparse·oracle·soft thresholding) · elastic net·bridge·fused lasso · LLM으로 능형·라쏘 돌릴 때 시험 주의 (Ch7 자료 배포)",
        topics=["AIC·BIC 유도: KL 발산, tr(J⁻¹I) → p, Laplace 근사 → p·log n", "제약 형태 ⟺ 벌점 형태(원과 마름모)", "능형회귀 해·해 경로·표준화", "1차원 능형 β̂ = y/(1+λ)", "라쏘: sparse vs dense, oracle property", "elastic net, 튜닝 파라미터 2개", "bridge 벌점", "soft thresholding", "grouping effect, generalized·fused lasso", "보건 응용과 LLM 시대의 메시지", "시험장 점검표(AI로 능형·라쏘)", "Chap7 Beyond Linearity 슬라이드(배포, 다음 차시)"],
        summary_pdf="8차시 시험대비 학습자료.pdf", full_pdf="8차시 시험대비 학습자료(전체판).pdf",
        # 20261022_강의노트.pdf 는 1~8차시 누적 노트북 → 8차시(2026-10-21) 부분만 추출
        notes=[("20261022_강의노트.pdf", {"start_marker": "2026년10월21일", "label": "누적 노트북 중 8차시 부분"})],
        slides=[("Ch7_BeyondLinear.pdf", "chap7_beyond_linearity", "ISLR Ch.7 (배포됨, 8차시에는 미진행 · 이미지 위주 → OCR)"),
                ("Ch7_Supplement_Proofs.pdf", "chap7_supplement_proofs", "Ch.7 보충 증명 (배포됨, 8차시에는 미진행)")],
        code=[],
        reuse=[("AIC_BIC_Derivations.pdf", "week06_20261007/slides/aic_bic_derivations.md"),
               ("ModelSelection_Regularization.pdf", "week06_20261007/slides/model_selection_regularization.md"),
               ("OtherPenaltyFunctions.pdf", "week06_20261007/slides/other_penalty_functions.md")],
        skip=[("20261022_강의노트.pdf (p.1~235)", "1~7차시 누적 부분은 각 주차의 강의노트와 동일하여 8차시 부분(p.236~)만 변환")],
    ),
]

# 2차시 논문읽기과제 폴더
PAPER_ASSIGNMENT = [
    # (원본 상대경로, 출력 파일명, 제목, 변환 방식)
    ("20260909_발표PPT_v2.1.pptx", "efron_presentation_final.md", "Efron(2020) 논문 발표 슬라이드(최종본, 발표자 노트 포함)", "pptx"),
    ("논문읽기과제/발제문_Efron_2020_source.html", "efron_presentation_brief.md", "Efron(2020) 발제문(1쪽)", "html"),
    ("논문읽기과제/Efron_2020_Prediction_Estimation_Attribution_요약본_3쪽_source.html", "efron_summary_3pages.md", "Efron(2020) 요약본(3쪽)", "html"),
    ("논문읽기과제/Efron_2020_Prediction_Estimation_Attribution_완전요약본_source.html", "efron_summary_full.md", "Efron(2020) 완전요약본(28쪽)", "html"),
    ("논문읽기과제/Efron_Prediction_Estimation_Attribution_한국어_요약.pdf", "efron_summary_korean_short.md", "Efron(2020) 한국어 요약(4쪽)", "pdf-doc"),
    ("논문읽기과제/Efron_예측_추정_귀속_논문_해설_요약.docx", "efron_guide_beginner.md", "Efron(2020) 입문자용 해설 요약(14쪽, 표 15개)", "docx"),
]

# ---------------------------------------------------------------------------
# 공통 유틸
# ---------------------------------------------------------------------------
_enc = None


def tokens(text: str) -> int:
    global _enc
    if _enc is None:
        import tiktoken
        _enc = tiktoken.get_encoding("o200k_base")
    return len(_enc.encode(text, disallowed_special=()))



def rel(path: Path) -> str:
    return str(path.relative_to(ROOT)).replace(os.sep, "/")


# ---------------------------------------------------------------------------
# PDF → Markdown (구현은 tools/pdfmd.py)
# ---------------------------------------------------------------------------
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pdfmd import pdf_to_md, join_lines as join_wrapped, joiner, clean_text, md_escape_cell, BULLET_RE, HANGUL  # noqa: E402


# ---------------------------------------------------------------------------
# PPTX / DOCX / HTML / 받아쓰기
# ---------------------------------------------------------------------------
def pptx_to_md(path: Path) -> tuple[str, dict]:
    from pptx import Presentation
    from pptx.util import Emu  # noqa: F401

    prs = Presentation(path)
    out = []
    for n, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = clean_text(slide.shapes.title.text_frame.text.replace("\n", " "))
        if not title:
            for sh in slide.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    title = clean_text(sh.text_frame.text.replace("\n", " "))[:80]
                    break
        out.append(f"\n## 슬라이드 {n}" + (f": {title}" if title else "") + "\n")
        shapes = sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
        for sh in shapes:
            if sh.has_text_frame:
                paras = []
                for p in sh.text_frame.paragraphs:
                    t = clean_text("".join(r.text for r in p.runs) or p.text)
                    if not t:
                        continue
                    if t == title and sh == slide.shapes.title:
                        continue
                    lvl = p.level or 0
                    paras.append(("  " * lvl + "- " + t) if (lvl > 0 or len(sh.text_frame.paragraphs) > 1) else t)
                if paras:
                    out.append("\n".join(paras) + "\n")
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = []
                for r_i, row in enumerate(sh.table.rows):
                    cells = [md_escape_cell(clean_text(c.text)) for c in row.cells]
                    rows.append("| " + " | ".join(cells) + " |")
                    if r_i == 0:
                        rows.append("|" + "---|" * len(cells))
                out.append("\n".join(rows) + "\n")
            if sh.shape_type == 13:
                out.append("_[그림]_\n")
        if slide.has_notes_slide:
            nt = clean_text(slide.notes_slide.notes_text_frame.text)
            if nt:
                notes_lines = [ln.strip() for ln in slide.notes_slide.notes_text_frame.text.splitlines() if ln.strip()]
                out.append("> **발표자 노트:** " + "\n> ".join(notes_lines) + "\n")
    md = re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n"
    return md, {"slides": len(prs.slides), "method": "pptx"}


def docx_to_md(path: Path) -> tuple[str, dict]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(path)
    out = []
    ntab = 0
    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, d)
            t = clean_text(p.text)
            if not t:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            m = re.match(r"Heading (\d)", style)
            if style == "Title":
                out.append("## " + t)
            elif m:
                out.append("#" * min(int(m.group(1)) + 1, 6) + " " + t)
            elif "List" in style or child.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr") is not None:
                out.append("- " + t)
            else:
                bold = all(r.bold for r in p.runs if r.text.strip()) and len(t) <= 80 and p.runs
                out.append(("**" + t + "**") if bold else t)
        elif tag == "tbl":
            tb = Table(child, d)
            ntab += 1
            rows = []
            for r_i, row in enumerate(tb.rows):
                cells = [md_escape_cell(clean_text(c.text)) for c in row.cells]
                rows.append("| " + " | ".join(cells) + " |")
                if r_i == 0:
                    rows.append("|" + "---|" * len(cells))
            out.append("\n".join(rows))
    md = "\n\n".join(out).strip() + "\n"
    return md, {"tables": ntab, "method": "docx"}


def html_to_md(path: Path) -> tuple[str, dict]:
    from bs4 import BeautifulSoup
    from markdownify import markdownify

    raw = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw, "html.parser")
    for t in soup(["style", "script", "img", "svg"]):
        t.decompose()
    md = markdownify(str(soup.body or soup), heading_style="ATX", bullets="-")
    md = html.unescape(md)
    md = re.sub(r"[ \t]+\n", "\n", md)
    md = re.sub(r"\n{3,}", "\n\n", md).strip() + "\n"
    # 문서 제목(h1)은 우리 헤더에서 쓰므로 한 단계 내림
    md = re.sub(r"^# ", "## ", md, flags=re.M)
    return md, {"method": "html"}


TS_RE = re.compile(r"^\[(\d{1,3}):(\d{2})\]\s*(.*)$")


def transcript_to_md(path: Path) -> tuple[str, dict]:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    out = []
    cur_block = -1
    last_ts = 0
    n = 0
    for ln in lines:
        m = TS_RE.match(ln.strip())
        if not m:
            if ln.strip():
                out.append(ln.strip())
            continue
        mm, ss, txt = int(m.group(1)), int(m.group(2)), clean_text(m.group(3))
        sec = mm * 60 + ss
        last_ts = max(last_ts, sec)
        block = sec // 300
        if block != cur_block:
            cur_block = block
            out.append(f"\n## {block*5:02d}:00 – {block*5+5:02d}:00\n")
        if txt:
            out.append(f"[{mm:02d}:{ss:02d}] {txt}  ")
            n += 1
    md = "\n".join(out).strip() + "\n"
    return md, {"lines": n, "duration": f"{last_ts//60}분 {last_ts%60}초", "method": "txt"}


# ---------------------------------------------------------------------------
# 분할 + 저장
# ---------------------------------------------------------------------------
def split_md(md: str, max_tokens: int) -> list[str]:
    if tokens(md) <= max_tokens:
        return [md]
    # 제목 경계로 세그먼트화 (## 우선, 부족하면 ###)
    for level in ("## ", "### ", "#### "):
        segs: list[str] = []
        cur: list[str] = []
        for ln in md.splitlines(keepends=True):
            if ln.startswith(level) and cur:
                segs.append("".join(cur))
                cur = []
            cur.append(ln)
        if cur:
            segs.append("".join(cur))
        if len(segs) > 1 and max(tokens(s) for s in segs) <= max_tokens:
            break
    # 여전히 큰 세그먼트는 빈 줄(문단) 경계로 쪼갬
    fine: list[str] = []
    for s in segs:
        if tokens(s) <= max_tokens:
            fine.append(s)
            continue
        paras = re.split(r"(\n\s*\n)", s)
        chunk = ""
        for p in paras:
            if tokens(chunk + p) > max_tokens and chunk.strip():
                fine.append(chunk)
                chunk = ""
            chunk += p
        if chunk.strip():
            fine.append(chunk)
    # 세그먼트를 max_tokens 이하로 묶기
    parts: list[str] = []
    cur_txt = ""
    for s in fine:
        if cur_txt and tokens(cur_txt + s) > max_tokens:
            parts.append(cur_txt)
            cur_txt = ""
        cur_txt += s
    if cur_txt.strip():
        parts.append(cur_txt)
    return [p.strip() + "\n" for p in parts if p.strip()]


EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
# 하이픈/점으로 구분된 한국 전화번호 형식만(OCR로 읽힌 표의 숫자열을 잘못 지우지 않도록 구분자 필수)
PHONE_RE = re.compile(r"(?<!\d)0\d{1,2}[-.]\d{3,4}[-.]\d{4}(?!\d)")


# 실명 치환 목록은 공개 소스에 두지 않고 private/redact.json 에서 읽는다.
#   형식: [["치환할 이름", "[발표자]"], ...]  (private/ 는 .gitignore 로 제외)
def _load_redact() -> list[tuple[str, str]]:
    f = PRIV / "redact.json"
    if f.exists():
        try:
            return [tuple(x) for x in json.loads(f.read_text(encoding="utf-8"))]
        except Exception as e:
            print("경고: private/redact.json 을 읽을 수 없음:", e)
    return []


REDACT = _load_redact()


def scrub_pii(text: str) -> str:
    """공개 저장소에 올라가면 안 되는 이메일·전화번호·동료 실명 제거(OCR/받아쓰기에 섞여 들어오는 경우 대비)."""
    text = EMAIL_RE.sub("[이메일 제거]", text)
    text = PHONE_RE.sub("[전화번호 제거]", text)
    for src, dst in REDACT:
        # 띄어쓰기 복원 과정에서 이름이 '홍 길 동'처럼 갈라질 수 있으므로 음절 사이 공백을 허용해 치환
        text = re.sub(r"\s*".join(map(re.escape, src)), dst, text)
    return text


def write_doc(md: str, out_path: Path, title: str, source_rel: str, method: str, week_readme: str,
              note: str = "", max_tokens: int = MAX_TOKENS, extra_meta: str = "") -> list[dict]:
    """md 를 분할하여 out_path(또는 out_path 의 _partNN 변형)로 저장하고 파일 목록(+토큰 수) 반환."""
    md = scrub_pii(md)
    parts = split_md(md, max_tokens - HEADER_RESERVE)  # 파트 머리말(제목·출처·링크) 토큰 여유
    total = len(parts)
    paths = []
    for k in range(total):
        if total == 1:
            paths.append(out_path)
        else:
            paths.append(out_path.with_name(f"{out_path.stem}_part{k+1:02d}{out_path.suffix}"))
    files = []
    for k, (p, body) in enumerate(zip(paths, parts)):
        p.parent.mkdir(parents=True, exist_ok=True)
        nav = []
        if total > 1:
            if k > 0:
                nav.append(f"[◀ 이전 부분]({paths[k-1].name})")
            if k < total - 1:
                nav.append(f"[다음 부분 ▶]({paths[k+1].name})")
        depth = len(p.relative_to(OUT).parts) - 1
        up = "../" * depth
        nav.append(f"[주차 인덱스]({os.path.relpath((OUT / week_readme), p.parent).replace(os.sep, '/')})")
        nav.append(f"[전체 인덱스]({up}../INDEX.md)")
        head = [
            "<!-- 자동 생성 파일: tools/convert_sources.py 가 원본에서 변환. 직접 수정하지 말고 원본/스크립트를 고치세요. -->",
            f"# {title}" + (f" ({k+1}/{total})" if total > 1 else ""),
            "",
            f"> 원본: `{source_rel}` · 변환 방식: {method} · 생성일: {TODAY}" + (f" · {extra_meta}" if extra_meta else ""),
        ]
        if note:
            head.append(f"> ⚠️ {note}")
        head.append("> " + " · ".join(nav))
        head.append("")
        head.append("---")
        head.append("")
        text = scrub_pii("\n".join(head) + body)  # 머리말의 원본 파일명에 든 실명도 치환
        p.write_text(text, encoding="utf-8")
        files.append({"path": rel(p), "tokens": tokens(text), "chars": len(text)})
    return files


# ---------------------------------------------------------------------------
# 주차 처리
# ---------------------------------------------------------------------------
def rec_id(txt_name: str) -> str:
    m = re.match(r"\d{8}-(\w+) 받아쓰기", txt_name)
    key = m.group(1) if m else "x"
    return f"rec{key}" if key.isdigit() else f"rec_{key}"


def process_week(w: dict, use_ocr: bool, max_tokens: int) -> dict:
    folder = SRC / w["folder"]
    wdir = OUT / w["slug"]
    for d in (wdir, PRIV_OUT / w["slug"]):
        if d.exists():
            shutil.rmtree(d)
    wdir.mkdir(parents=True)
    week_readme = f"{w['slug']}/README.md"
    entries: list[dict] = []
    log = lambda *a: print(f"[{w['slug']}]", *a, flush=True)

    def add(category, title, files, source, method, note="", stats=None):
        entries.append(dict(category=category, title=title, files=files, source=source, method=method,
                            note=note, stats=stats or {}))

    # 시험대비 학습자료(요약본)
    if w.get("summary_pdf"):
        src = folder / w["summary_pdf"]
        log("요약본", src.name)
        md, st = pdf_to_md(src, "doc", use_ocr=use_ocr)
        files = write_doc(md, wdir / "study_guide_summary.md", f"{w['n']}차시 시험대비 학습자료 (요약본)", rel(src),
                          st["method"], week_readme, max_tokens=max_tokens,
                          extra_meta=f"원본 {st['pages']}쪽")
        add("study_guide_summary", f"{w['n']}차시 시험대비 학습자료 (요약본)", files, rel(src), st["method"], stats=st)

    # 시험대비 학습자료(전체판)
    if w.get("full_pdf"):
        src = folder / w["full_pdf"]
        log("전체판", src.name)
        md, st = pdf_to_md(src, "doc", use_ocr=use_ocr)
        files = write_doc(md, wdir / "study_guide_full" / "study_guide_full.md",
                          f"{w['n']}차시 시험대비 학습자료 (전체판)", rel(src), st["method"], week_readme,
                          note="전체판의 슬라이드 이미지는 텍스트로 옮기지 못했습니다. 이미지 위치는 _[그림/슬라이드 이미지]_ 로 표시했고, 슬라이드 내용은 '슬라이드 원문' 박스와 slides/ 폴더를 참고하세요.",
                          max_tokens=max_tokens, extra_meta=f"원본 {st['pages']}쪽")
        add("study_guide_full", f"{w['n']}차시 시험대비 학습자료 (전체판)", files, rel(src), st["method"], stats=st)

    # 강의노트
    for name, opt in w.get("notes", []):
        src = folder / name
        import pymupdf
        pages = None
        extra = ""
        if opt and opt.get("start_marker"):
            d = pymupdf.open(src)
            start = next((i for i in range(len(d)) if opt["start_marker"] in d[i].get_text()), None)
            if start is None:
                log("경고: start_marker 를 찾지 못함", name)
                start = 0
            pages = list(range(start, len(d)))
            extra = f"{opt.get('label','')} (원본 p.{start+1}~{len(d)})"
        log("강의노트", name, extra)
        md, st = pdf_to_md(src, "notes", pages=pages, use_ocr=use_ocr)
        files = write_doc(md, wdir / "lecture_notes.md", f"{w['n']}차시 강의노트 (OneNote 내보내기)", rel(src), st["method"],
                          week_readme,
                          note="OneNote 필기 PDF에서 텍스트 층을 추출했고, 슬라이드 캡처만 있는 페이지는 OCR로 읽었습니다. 손글씨·수식은 누락될 수 있습니다. 교수님 말씀 원문은 전체판의 '부록 A 강의노트 원문'이 더 정확합니다.",
                          max_tokens=max_tokens, extra_meta=(extra or f"원본 {st['pages']}쪽") + f" · OCR {st['ocr_pages']}쪽")
        add("lecture_notes", f"{w['n']}차시 강의노트", files, rel(src), st["method"], note=extra, stats=st)

    # 슬라이드
    for name, slug, desc in w.get("slides", []):
        src = folder / name
        if "*" in name:  # 파일명에 실명이 든 원본은 글롭 패턴으로 지정
            matches = sorted(folder.glob(name))
            if not matches:
                log("경고: 원본 없음", name)
                continue
            src = matches[0]
        log("슬라이드", src.name)
        label = desc if "*" in name else f"{Path(name).stem} — {desc}"
        md, st = pdf_to_md(src, "slides", use_ocr=use_ocr)
        files = write_doc(md, wdir / "slides" / f"{slug}.md", f"슬라이드: {label}", rel(src), st["method"],
                          week_readme, max_tokens=max_tokens,
                          extra_meta=f"{st['pages']}장" + (f" · OCR {st['ocr_pages']}장" if st["ocr_pages"] else ""))
        add("slides", label, files, rel(src), st["method"], stats=st)

    # 받아쓰기
    for src in sorted(folder.glob("*받아쓰기(자동).txt")):
        rid = rec_id(src.name)
        log("받아쓰기", src.name)
        md, st = transcript_to_md(src)
        label = "논문/기사 녹음" if "article" in rid else f"녹음 {rid[3:]}"
        files = write_doc(md, wdir / "transcripts" / f"{rid}.md", f"{w['n']}차시 {label} 받아쓰기 (Whisper 자동)", rel(src), "txt",
                          week_readme,
                          note="Whisper(large-v3-turbo) 자동 받아쓰기라 오인식이 많습니다. 용어 검색·맥락 파악용으로만 쓰고, 정확한 표현은 학습자료의 녹음 인용 박스와 대조하세요.",
                          max_tokens=max_tokens, extra_meta=f"길이 {st['duration']} · {st['lines']}줄")
        add("transcripts", f"{label} ({st['duration']})", files, rel(src), "txt", stats=st)

    # 코드
    for name, out_name, desc in w.get("code", []):
        src = folder / name
        (wdir / "code").mkdir(exist_ok=True)
        body = src.read_text(encoding="utf-8", errors="replace").replace("\r\n", "\n").strip() + "\n"
        header = f"# {desc}\n# 원본: {rel(src)} (수업 시간에 배포된 R 코드) · 변환일 {TODAY}\n\n"
        p = wdir / "code" / out_name
        p.write_text(header + body, encoding="utf-8")
        add("code", desc, [{"path": rel(p), "tokens": tokens(header + body), "chars": len(header + body)}], rel(src), "copy")

    # 논문읽기과제(2차시)
    if w.get("paper_assignment"):
        for name, out_name, title, how in PAPER_ASSIGNMENT:
            src = folder / name
            log("논문과제", name)
            if how == "pptx":
                md, st = pptx_to_md(src)
            elif how == "html":
                md, st = html_to_md(src)
            elif how == "docx":
                md, st = docx_to_md(src)
            else:
                md, st = pdf_to_md(src, "doc", use_ocr=use_ocr)
            files = write_doc(md, wdir / "paper_assignment" / out_name, title, rel(src), st["method"], week_readme,
                              max_tokens=max_tokens)
            add("paper_assignment", title, files, rel(src), st["method"], stats=st)

    # 원본 폴더 통계(미변환 자료)
    photos = sorted(glob.glob(str(folder / "Quick Share" / "*.jpg")))
    audio = sorted(glob.glob(str(folder / "*.wma")) + glob.glob(str(folder / "*.m4a")) + glob.glob(str(folder / "Quick Share" / "*.m4a")))
    size = sum(os.path.getsize(f) for f in glob.glob(str(folder / "**" / "*"), recursive=True) if os.path.isfile(f))
    return dict(n=w["n"], slug=w["slug"], folder=w["folder"], date=w["date"], title=w["title"], topics=w["topics"],
                entries=entries, reuse=w.get("reuse", []), skip=w.get("skip", []),
                photos=len(photos), audio=len(audio), src_bytes=size)


def process_common() -> list[dict]:
    cdir = PRIV_OUT / "common"
    cdir.mkdir(parents=True, exist_ok=True)
    entries = []
    src = SRC / "학습자료_생성_프롬프트.md"
    body = re.sub(r"C:\\Users\\[^\\]+\\", r"C:\\Users\\<사용자>\\", src.read_text(encoding="utf-8"))
    text = ("<!-- 원본: machine_learning_for_health_science/학습자료_생성_프롬프트.md 를 복사(로컬 사용자 이름만 가림). -->\n"
            + body.rstrip() + "\n")
    p = cdir / "study_material_prompt.md"
    p.write_text(text, encoding="utf-8")
    entries.append(dict(category="common", title="주차별 시험대비 학습자료 생성 프롬프트", files=[{"path": rel(p), "tokens": tokens(text), "chars": len(text)}],
                        source=rel(src), method="copy", note="", stats={}))

    refs = f"""# 교재·논문 참고 링크

> 저작권이 있는 교재와 학술지 논문 원문은 이 저장소에 넣지 않았습니다(원본 폴더에만 보관, git 제외). 아래 공식 링크를 이용하세요. 생성일 {TODAY}

| 자료 | 어디서 쓰였나 | 공식 링크 |
|---|---|---|
| James, Witten, Hastie, Tibshirani. *An Introduction to Statistical Learning with Applications in R* (ISLR, 1판) — 주교재 | 전 주차. 학습자료의 노란 박스 `[교재 보충 · ISLR p.N]` 쪽수는 1판 기준 | <https://www.statlearning.com/> (저자 공식 무료 PDF) |
| Efron, B. (2020). *Prediction, Estimation, and Attribution.* International Statistical Review 88(S1), S28–S59 (JASA 115(530), 636–655 재수록) | 1~2차시 읽기 과제, 2차시 논문 발표 | <https://doi.org/10.1111/insr.12409> · <https://doi.org/10.1080/01621459.2020.1762613> |
| Shmueli, G. (2010). *To Explain or to Predict?* Statistical Science 25(3), 289–310 | 2차시 동료 학생 발표 | <https://doi.org/10.1214/10-STS330> |
| Belkin, Hsu, Ma, Mandal (2019). *Reconciling modern machine-learning practice and the classical bias–variance trade-off.* PNAS 116(32), 15849–15854 | 3차시 학생 발표(Double descent) | <https://doi.org/10.1073/pnas.1903070116> |
| Prince, S. *Understanding Deep Learning* — 부교재(딥러닝 파트 예정) | 1차시 과목 소개에서 언급 | <https://udlbook.github.io/udlbook/> |
| Hastie, Tibshirani, Friedman. *The Elements of Statistical Learning* (ESL) — 참고서 | 1차시 과목 소개에서 언급 | <https://hastie.su.domains/ElemStatLearn/> |
"""
    p = cdir / "references.md"
    p.write_text(refs, encoding="utf-8")
    entries.append(dict(category="common", title="교재·논문 참고 링크", files=[{"path": rel(p), "tokens": tokens(refs), "chars": len(refs)}],
                        source="(직접 작성)", method="manual", note="", stats={}))
    return entries


# ---------------------------------------------------------------------------
# 인덱스 생성
# ---------------------------------------------------------------------------
CATEGORY_LABEL = {
    "study_guide_summary": "시험대비 학습자료(요약본)",
    "study_guide_full": "시험대비 학습자료(전체판)",
    "lecture_notes": "강의노트",
    "slides": "슬라이드",
    "transcripts": "녹음 받아쓰기",
    "code": "R 코드",
    "paper_assignment": "논문읽기과제",
    "common": "공통",
    "compiled": "모아보기",
}
CATEGORY_ORDER = ["study_guide_summary", "study_guide_full", "lecture_notes", "slides", "paper_assignment", "transcripts", "code"]


def is_public(entry: dict) -> bool:
    return entry["category"] in PUBLIC_CATEGORIES


def link(path_rel_root: str, from_dir_rel_root: str, text: str | None = None) -> str:
    target = os.path.relpath(path_rel_root, from_dir_rel_root or ".").replace(os.sep, "/")
    return f"[{text or Path(path_rel_root).name}]({target})"


def files_cell(files: list[dict], from_dir: str, index_path: str | None = None) -> str:
    if len(files) == 1:
        return link(files[0]["path"], from_dir)
    parts = " · ".join(link(f["path"], from_dir, f"{k+1}부") for k, f in enumerate(files))
    if index_path:
        return link(index_path, from_dir, "목차") + " · " + parts
    return parts


# ---------------------------------------------------------------------------
# 공개 범위 분리: 공개 카테고리 문서는 docs/, 그 외는 private/docs/ (같은 하위 구조)
# ---------------------------------------------------------------------------
def _rebase(path_rel: str, to_root: Path) -> Path:
    parts = Path(path_rel).parts
    if parts[:2] == ("private", "docs"):
        return to_root.joinpath(*parts[2:])
    if parts[:1] == ("docs",):
        return to_root.joinpath(*parts[1:])
    return ROOT / path_rel


def partition_scope(weeks: list[dict]):
    """주차 manifest 의 파일들을 공개/비공개 루트로 옮기고 경로를 갱신한다(이미 제자리면 그대로)."""
    for w in weeks:
        for e in w["entries"]:
            e["public"] = is_public(e)
            root = OUT if e["public"] else PRIV_OUT
            for f in e["files"]:
                cur = ROOT / f["path"]
                dest = _rebase(f["path"], root)
                if cur.resolve() != dest.resolve():
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    if cur.exists():
                        shutil.move(str(cur), str(dest))
                    f["path"] = rel(dest)
            if e.get("index"):
                old = ROOT / e["index"]
                if old.exists() and _rebase(e["index"], root).resolve() != old.resolve():
                    old.unlink()
                e.pop("index", None)
    for root in (OUT, PRIV_OUT):
        if root.exists():
            for d in sorted((p for p in root.rglob("*") if p.is_dir()), key=lambda p: -len(p.parts)):
                try:
                    d.rmdir()  # 빈 폴더만 지워짐
                except OSError:
                    pass


# ---------------------------------------------------------------------------
# 분할된 문서의 목차 파일
# ---------------------------------------------------------------------------
def _part_headings(path: str, limit: int = 14) -> str:
    heads = []
    for ln in (ROOT / path).read_text(encoding="utf-8").splitlines():
        m = re.match(r"^(#{2,3})\s+(.*)", ln)
        if m:
            heads.append(m.group(2).strip())
    if not heads:
        for ln in (ROOT / path).read_text(encoding="utf-8").splitlines():
            m = re.match(r"^#{4}\s+(.*)", ln)
            if m:
                heads.append(m.group(1).strip())
    shown = heads[:limit]
    txt = " · ".join(h if len(h) <= 60 else h[:57] + "…" for h in shown)
    if len(heads) > limit:
        txt += f" … (+{len(heads) - limit})"
    return md_escape_cell(txt) or "—"


def write_stubs(entries: list[dict]):
    for e in entries:
        if len(e["files"]) < 2:
            e.pop("index", None)
            continue
        first = e["files"][0]["path"]
        stub_rel = re.sub(r"_part\d+(\.md)$", r"\1", first)
        if stub_rel == first:
            continue
        public = e.get("public", True)
        rows = []
        for k, f in enumerate(e["files"], 1):
            extra = f" | `{RAW_BASE}{f['path']}` |" if public else " |"
            rows.append(f"| [{k}부]({Path(f['path']).name}) | {_part_headings(f['path'])} | {f['tokens']:,}{extra}")
        root = OUT if public else PRIV_OUT
        depth = len(Path(stub_rel).relative_to(rel(root)).parts) - 1
        index_link = ("../" * depth) + "../INDEX.md"
        head_cols = "| 부 | 포함된 절(제목) | tokens | raw URL |\n|---|---|---|---|" if public else "| 부 | 포함된 절(제목) | tokens |\n|---|---|---|"
        text = "\n".join([
            "<!-- 자동 생성 목차 파일: tools/convert_sources.py -->",
            f"# {e['title']} — 목차",
            "",
            f"> 이 문서는 {len(e['files'])}부로 나뉘어 있습니다(파일당 최대 {MAX_TOKENS:,} tokens). 아래 표에서 필요한 절이 들어 있는 부를 고르세요. "
            + ("ChatGPT에는 raw URL 열의 주소를 그대로 붙여 넣으면 됩니다. " if public else "비공개 문서(로컬 전용)입니다. ")
            + f"원본: `{e['source']}` · 생성일 {TODAY} · [전체 인덱스]({index_link})",
            "",
            head_cols,
            *rows,
            "",
        ])
        (ROOT / stub_rel).parent.mkdir(parents=True, exist_ok=True)
        (ROOT / stub_rel).write_text(scrub_pii(text), encoding="utf-8")
        e["index"] = stub_rel


# ---------------------------------------------------------------------------
# 모아보기: 각 차시 학습자료의 같은 성격 섹션을 한 파일씩으로 합침 → 학기 전체를 한 URL로 검색 (공개)
# ---------------------------------------------------------------------------
COMPILED = [
    ("overviews.md", "전 차시 강의 개요 모음", ["강의 개요"], "각 차시 학습자료의 '강의 개요' 절만 모았습니다. 학기 흐름을 한눈에 볼 때."),
    ("summary_tables.md", "핵심 개념 한 장 요약표 모음", ["핵심 개념 한 장 요약표"], "질문/한 줄 답/출처 형식의 요약표. 시험 직전 복습용."),
    ("glossary.md", "용어 정리 모음 (학기 용어집)", ["용어 정리", "용어표"], "각 차시의 용어 정리 절을 모은 학기 용어집."),
    ("key_points.md", "교수님이 강조한 포인트 모음", ["교수님이 강조한 포인트"], "녹음·강의노트에서 교수님이 반복·강조한 내용(시험 관련 발언 포함)."),
    ("exam_questions.md", "예상 시험 문제와 모범 답안 모음", ["예상 시험 문제"], "객관식·단답·서술형 예상문제와 답. 근거 출처 표시 [노트]/[녹음]/[교재]/[추가]."),
    ("checklists.md", "셀프 체크리스트 모음", ["셀프 체크리스트"], "☐ 항목으로 된 자가 점검표."),
]
HEADING_RE = re.compile(r"^(#{2,4})\s+(.*)$")


def _strip_part_header(text: str) -> str:
    idx = text.find("\n---\n")
    return text[idx + 5:] if idx >= 0 else text


def _joined_doc(files: list[dict]) -> str:
    return "\n".join(_strip_part_header((ROOT / f["path"]).read_text(encoding="utf-8")) for f in files)


def _extract_sections(md: str, keywords: list[str]) -> list[str]:
    lines = md.splitlines()
    out = []
    i = 0
    while i < len(lines):
        m = HEADING_RE.match(lines[i])
        if m and any(k in m.group(2) for k in keywords) and "부록" not in m.group(2):
            level = len(m.group(1))
            j = i + 1
            while j < len(lines):
                m2 = HEADING_RE.match(lines[j])
                if m2 and len(m2.group(1)) <= level:
                    break
                j += 1
            body = "\n".join(lines[i:j]).strip()
            body = re.sub(r"<!-- page \d+ -->\n?", "", body)
            body = re.sub(r"^#{2,4}\s+", "### ", body, count=1)
            out.append(body)
            i = j
        else:
            i += 1
    return out


def compile_sections(weeks: list[dict], max_tokens: int) -> list[dict]:
    cdir = OUT / "compiled"
    if cdir.exists():
        shutil.rmtree(cdir)
    cdir.mkdir(parents=True)
    entries = []
    for fname, title, keywords, desc in COMPILED:
        chunks = [f"> {desc} 각 차시 학습자료(요약본, 없으면 전체판)에서 자동 추출 · 생성일 {TODAY} · [전체 인덱스](../../INDEX.md)\n"]
        found_weeks = 0
        for w in weeks:
            src_entries = [e for e in w["entries"] if e["category"] == "study_guide_summary"] or \
                          [e for e in w["entries"] if e["category"] == "study_guide_full"]
            if not src_entries:
                continue
            src = src_entries[0]
            secs = _extract_sections(_joined_doc(src["files"]), keywords)
            if not secs:
                continue
            found_weeks += 1
            if src.get("public", is_public(src)):
                origin = f"[{Path(src['files'][0]['path']).name}]({os.path.relpath(src['files'][0]['path'], 'docs/compiled')})"
            else:
                origin = f"{CATEGORY_LABEL[src['category']]}(비공개, 로컬 private/ 폴더)"
            chunks.append(f"\n## {w['n']}차시 ({w['date']}) — {w['title']}\n\n_출처: {origin} · [주차 인덱스](../{w['slug']}/README.md)_\n")
            chunks.extend(sec + "\n" for sec in secs)
        if found_weeks == 0:
            continue
        files = write_doc("\n".join(chunks), cdir / fname, title, "docs/week*/study_guide_*.md", "compiled", "../INDEX.md",
                          max_tokens=max_tokens, extra_meta=f"{found_weeks}개 차시")
        entries.append(dict(category="compiled", title=title, files=files, source="각 차시 학습자료", method="compiled",
                            note=desc, stats={"weeks": found_weeks}, public=True))
    return entries


# ---------------------------------------------------------------------------
# 주차 README (공개용 docs/… 와 전체 목록용 private/docs/…)
# ---------------------------------------------------------------------------
def _entries_table(entries: list[dict], from_dir: str) -> list[str]:
    out = ["| 유형 | 문서 | 파트 | 분량(tokens) | 원본 | 비고 |", "|---|---|---|---|---|---|"]
    for cat in CATEGORY_ORDER:
        for e in [e for e in entries if e["category"] == cat]:
            tot = sum(f["tokens"] for f in e["files"])
            note = e.get("note") or ""
            if e["stats"].get("ocr_pages"):
                note = (note + " · " if note else "") + f"OCR {e['stats']['ocr_pages']}쪽"
            out.append(f"| {CATEGORY_LABEL[cat]} | {e['title']} | {files_cell(e['files'], from_dir, e.get('index'))} | {tot:,} | `{Path(e['source']).name}` | {note} |")
    return out


def write_week_readme(w: dict):
    pub = [e for e in w["entries"] if e.get("public")]
    priv = [e for e in w["entries"] if not e.get("public")]
    head = f"# {w['n']}차시 ({w['date']}) — {w['title']}"
    kw = " · ".join(f"`{t}`" for t in w["topics"])

    # 공개용
    wdir = f"docs/{w['slug']}"
    out = [head, "",
           f"> 변환 생성일 {TODAY} · [전체 인덱스](../../INDEX.md) · [주제별 찾기](../../TOPICS.md) · [모아보기](../compiled/)", "",
           "## 이번 차시 키워드", "", kw, "", "## 공개 문서 (시험대비 학습자료 요약본)", ""]
    if pub:
        out += _entries_table(pub, wdir)
    else:
        out.append("이 차시에는 요약본이 없습니다(전체판만 작성됨). 핵심 내용은 [모아보기](../compiled/)의 각 문서에서 `" + f"{w['n']}차시" + "` 절을 보세요.")
    counts = Counter(CATEGORY_LABEL[e["category"]] for e in priv)
    out += ["", "## 비공개 자료", "",
            "전체판·강의노트·슬라이드·녹음 받아쓰기·논문과제·코드는 저작권과 개인정보 때문에 공개하지 않습니다"
            + (f" (로컬 `private/` 폴더에 {', '.join(f'{k} {v}건' for k, v in counts.items())})." if counts else "."), ""]
    (OUT / wdir.split("/", 1)[1]).mkdir(parents=True, exist_ok=True)
    (ROOT / wdir / "README.md").write_text(scrub_pii("\n".join(out) + "\n"), encoding="utf-8")

    # 비공개(전체 목록)용
    pdir = f"private/docs/{w['slug']}"
    out = [head + " — 전체 목록 (비공개)", "",
           f"> 원본 폴더 `machine_learning_for_health_science/{w['folder']}/` ({w['src_bytes']/1e6:,.0f} MB, 사진 {w['photos']}장, 녹음 {w['audio']}개) · 변환 생성일 {TODAY} · [전체 인덱스(비공개)](../../INDEX.md) · [공개 README](../../../{wdir}/README.md)", "",
           "## 이번 차시 키워드", "", kw, "", "## 변환된 문서 (공개 + 비공개)", ""]
    out += _entries_table(w["entries"], pdir)
    if w.get("reuse"):
        out += ["", "## 다른 주차에서 이미 변환된 자료 (이 폴더에도 같은 원본이 있음)", "", "| 원본 | 변환 문서 |", "|---|---|"]
        for name, target in w["reuse"]:
            tgt = next((f["path"] for ww in CURRENT_WEEKS for e in ww["entries"] for f in e["files"][:1]
                        if f["path"].endswith("/" + target.split("/", 1)[1])), None)
            out.append(f"| `{name}` | {link(tgt, pdir) if tgt else target} |")
    out += ["", "## 변환하지 않은 원본", "", "| 원본 | 사유 |", "|---|---|"]
    if w["audio"]:
        out.append(f"| 녹음 파일 {w['audio']}개 (.wma/.m4a) | 용량·저작권. 내용은 `transcripts/` 받아쓰기로 대체 |")
    if w["photos"]:
        out.append(f"| 칠판 사진 {w['photos']}장 (Quick Share/*.jpg) | 손글씨 사진은 OCR 신뢰도가 낮아 보류. 전체판 학습자료가 촬영 시각 기준으로 사진을 배치했으므로 원본 PDF 참고 |")
    for name, why in w.get("skip", []):
        out.append(f"| `{name}` | {why} |")
    out.append("| `.fablize/` | 학습자료 생성 도구의 작업 기록(goals.json, ledger.jsonl) — 학습 내용 아님 |")
    (ROOT / pdir).mkdir(parents=True, exist_ok=True)
    (ROOT / pdir / "README.md").write_text(scrub_pii("\n".join(out) + "\n"), encoding="utf-8")


CURRENT_WEEKS: list[dict] = []


# ---------------------------------------------------------------------------
# 전체 인덱스 (공개 INDEX.md / 비공개 private/INDEX.md), llms.txt
# ---------------------------------------------------------------------------
def write_index(weeks: list[dict], common: list[dict], compiled: list[dict]):
    # ---- 공개 ----
    pub_files = [f for w in weeks for e in w["entries"] if e.get("public") for f in e["files"]] + [f for e in compiled for f in e["files"]]
    out = ["# 전체 문서 인덱스 — 보건학을 위한 기계학습 (2026-2)", "",
           f"> 자동 생성 {TODAY} (`python3 tools/convert_sources.py --index-only`) · 공개 문서 {len(pub_files)}개 · 총 {sum(f['tokens'] for f in pub_files):,} tokens · "
           f"파일 하나는 최대 {MAX_TOKENS:,} tokens(ChatGPT 웹에서 한 번에 읽히는 크기) · 공개 범위: 시험대비 학습자료 **요약본**과 **모아보기**만 "
           f"(전체판·강의노트·슬라이드·받아쓰기·논문과제는 비공개)", "",
           "## 0. 모아보기 (학기 전체를 한 파일로)", "", "| 문서 | 내용 | 파트 |", "|---|---|---|"]
    for e in compiled:
        out.append(f"| **{e['title']}** | {e['note']} | {files_cell(e['files'], '', e.get('index'))} |")
    out += ["", "## 1. 주차별 요약본", "", "| 차시 | 날짜 | 주제 | 요약본 | tokens |", "|---|---|---|---|---|"]
    for w in weeks:
        es = [e for e in w["entries"] if e.get("public")]
        if es:
            e = es[0]
            cell = link(e.get("index") or e["files"][0]["path"], "", f"{len(e['files'])}부" if len(e["files"]) > 1 else "요약본")
            tok = f"{sum(f['tokens'] for f in e['files']):,}"
        else:
            cell, tok = "요약본 없음 (모아보기 참고)", "—"
        out.append(f"| [{w['n']}차시](docs/{w['slug']}/README.md) | {w['date']} | {w['title']} | {cell} | {tok} |")
    out += ["", "## 2. 주차별 상세 (공개 파일)", ""]
    for w in weeks:
        out += [f"### {w['n']}차시 ({w['date']}) — {w['title']}", "", f"키워드: {' · '.join(w['topics'])}  ",
                f"주차 인덱스: [docs/{w['slug']}/README.md](docs/{w['slug']}/README.md)", ""]
        es = [e for e in w["entries"] if e.get("public")]
        if es:
            out += ["| 문서 | 파일 | tokens |", "|---|---|---|"]
            for e in es:
                for f in e["files"]:
                    out.append(f"| {e['title']} | {link(f['path'], '')} | {f['tokens']:,} |")
        else:
            out.append("공개 파일 없음(요약본 미작성). 모아보기 문서의 해당 차시 절을 보세요.")
        out.append("")
    ex = next((f["path"] for e in compiled for f in e["files"]), "docs/compiled/glossary.md")
    out += ["## 3. ChatGPT에 붙여 넣을 raw URL 규칙", "", f"`{RAW_BASE}` + 위 표의 파일 경로. 예:", "", f"```\n{RAW_BASE}{ex}\n```", ""]
    (ROOT / "INDEX.md").write_text(scrub_pii("\n".join(out) + "\n"), encoding="utf-8")

    # ---- 비공개(전체) ----
    all_files = [f for w in weeks for e in w["entries"] for f in e["files"]]
    out = ["# 전체 문서 인덱스 (비공개 · 로컬 전용) — 보건학을 위한 기계학습 (2026-2)", "",
           f"> 자동 생성 {TODAY} · 주차 문서 {len(all_files)}개({sum(f['tokens'] for f in all_files):,} tokens) + 모아보기 + 공통. "
           f"이 폴더(`private/`)는 .gitignore 로 제외되어 GitHub에 올라가지 않습니다. 공개 인덱스: [../INDEX.md](../INDEX.md)", "",
           "## 0. 모아보기 (공개)", "", "| 문서 | 파트 |", "|---|---|"]
    for e in compiled:
        out.append(f"| **{e['title']}** | {files_cell(e['files'], 'private', e.get('index'))} |")
    out += ["", "## 1. 공통 자료", "", "| 문서 | 내용 |", "|---|---|"]
    for e in common:
        out.append(f"| {link(e['files'][0]['path'], 'private')} | {e['title']} |")
    out += ["", "## 2. 주차 × 자료 유형", "", "| 차시 | 날짜 | 주제 | 요약본 | 전체판 | 강의노트 | 슬라이드 | 받아쓰기 | 기타 |", "|---|---|---|---|---|---|---|---|---|"]
    for w in weeks:
        def cnt(cat):
            es = [e for e in w["entries"] if e["category"] == cat]
            if not es:
                return "—"
            nfiles = sum(len(e["files"]) for e in es)
            first = es[0].get("index") or es[0]["files"][0]["path"]
            label = f"{len(es)}개" if cat in ("slides", "transcripts") else (f"{nfiles}부" if nfiles > 1 else "1")
            return link(first, "private", label)
        extras = []
        if any(e["category"] == "paper_assignment" for e in w["entries"]):
            extras.append(link(f"private/docs/{w['slug']}/paper_assignment", "private", "논문과제"))
        if any(e["category"] == "code" for e in w["entries"]):
            extras.append(link(f"private/docs/{w['slug']}/code", "private", "R코드"))
        out.append(f"| [{w['n']}차시](docs/{w['slug']}/README.md) | {w['date']} | {w['title']} | {cnt('study_guide_summary')} | {cnt('study_guide_full')} | {cnt('lecture_notes')} | {cnt('slides')} | {cnt('transcripts')} | {' · '.join(extras) or '—'} |")
    out += ["", "## 3. 주차별 상세 목록", ""]
    for w in weeks:
        out += [f"### {w['n']}차시 ({w['date']}) — {w['title']}", "", "| 유형 | 문서 | 파일 | tokens | 공개 |", "|---|---|---|---|---|"]
        for cat in CATEGORY_ORDER:
            for e in [e for e in w["entries"] if e["category"] == cat]:
                for f in e["files"]:
                    out.append(f"| {CATEGORY_LABEL[cat]} | {e['title']} | {link(f['path'], 'private')} | {f['tokens']:,} | {'공개' if e.get('public') else '비공개'} |")
        out.append("")
    PRIV.mkdir(parents=True, exist_ok=True)
    (PRIV / "INDEX.md").write_text(scrub_pii("\n".join(out) + "\n"), encoding="utf-8")


def write_llms_txt(weeks: list[dict], compiled: list[dict]):
    out = ["# 보건학을 위한 기계학습 (서울대 보건대학원, 2026-2) 학습 문서 모음", "",
           f"> {weeks[-1]['n'] if weeks else 0}차시까지의 시험대비 학습자료 요약본과 학기 모아보기(용어집·예상문제·강조 포인트·체크리스트·요약표·강의 개요)를 Markdown으로 정리한 저장소. "
           f"생성 {TODAY}. 파일당 최대 {MAX_TOKENS:,} tokens. 인덱스: {RAW_BASE}INDEX.md , 주제별: {RAW_BASE}TOPICS.md", "",
           "## 모아보기 (학기 전체)", ""]
    for e in compiled:
        for k, f in enumerate(e["files"]):
            suffix = f" ({k+1}/{len(e['files'])})" if len(e["files"]) > 1 else ""
            out.append(f"- [{e['title']}{suffix}]({RAW_BASE}{f['path']})")
    for w in weeks:
        out += ["", f"## {w['n']}차시 ({w['date']}) {w['title']}", "", f"- [주차 인덱스]({RAW_BASE}docs/{w['slug']}/README.md)"]
        for e in [e for e in w["entries"] if e.get("public")]:
            for k, f in enumerate(e["files"]):
                suffix = f" ({k+1}/{len(e['files'])})" if len(e["files"]) > 1 else ""
                out.append(f"- [{e['title']}{suffix}]({RAW_BASE}{f['path']})")
    (ROOT / "llms.txt").write_text(scrub_pii("\n".join(out) + "\n"), encoding="utf-8")


# ---------------------------------------------------------------------------
# 정리: 예전 실행이 남긴 파일(옛 목차 파일 등) 제거, 이름·연락처 재치환
# ---------------------------------------------------------------------------
def cleanup_trees(weeks: list[dict], common: list[dict], compiled: list[dict]):
    expected = set()
    for e in [e for w in weeks for e in w["entries"]] + list(common) + list(compiled):
        expected.update(f["path"] for f in e["files"])
        if e.get("index"):
            expected.add(e["index"])
    keep_names = {"README.md", "manifest.json", "week_manifest.json", "INDEX.md"}
    removed = 0
    for root in (OUT, PRIV_OUT):
        if not root.exists():
            continue
        for f in list(root.rglob("*")):
            if f.is_file() and rel(f) not in expected and f.name not in keep_names:
                f.unlink()
                removed += 1
        for d in sorted((p for p in root.rglob("*") if p.is_dir()), key=lambda p: -len(p.parts)):
            try:
                d.rmdir()
            except OSError:
                pass
    if removed:
        print(f"정리: 오래된 파일 {removed}개 삭제")


def rescrub_all():
    """REDACT/연락처 규칙이 바뀐 뒤에도 다시 변환하지 않고 기존 문서에 치환을 다시 적용."""
    changed = 0
    for root in (OUT, PRIV):
        if not root.exists():
            continue
        for f in root.rglob("*"):
            if f.is_file() and f.suffix in (".md", ".json", ".txt", ".R"):
                t = f.read_text(encoding="utf-8", errors="replace")
                t2 = scrub_pii(t)
                if t2 != t:
                    f.write_text(t2, encoding="utf-8")
                    changed += 1
    if changed:
        print(f"재치환: 파일 {changed}개 수정")


# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", help="예: week03 또는 week01,week02 (slug 앞부분, 쉼표로 여러 개)")
    ap.add_argument("--index-only", action="store_true")
    ap.add_argument("--no-ocr", action="store_true")
    ap.add_argument("--max-tokens", type=int, default=MAX_TOKENS)
    args = ap.parse_args()

    OUT.mkdir(exist_ok=True)
    PRIV_OUT.mkdir(parents=True, exist_ok=True)
    if not args.index_only:
        if not SRC.exists():
            sys.exit(f"원본 폴더가 없습니다: {SRC}")
        wanted = [x.strip() for x in args.only.split(",")] if args.only else None
        for w in WEEKS:
            if wanted and not any(w["slug"].startswith(x) for x in wanted):
                continue
            res = process_week(w, use_ocr=not args.no_ocr, max_tokens=args.max_tokens)
            (PRIV_OUT / w["slug"]).mkdir(parents=True, exist_ok=True)
            (PRIV_OUT / w["slug"] / "week_manifest.json").write_text(scrub_pii(json.dumps(res, ensure_ascii=False, indent=1)), encoding="utf-8")

    # 예전 위치(docs/)에 남은 주차 manifest·공통 자료 정리
    for w in WEEKS:
        legacy = OUT / w["slug"] / "week_manifest.json"
        if legacy.exists():
            (PRIV_OUT / w["slug"]).mkdir(parents=True, exist_ok=True)
            shutil.move(str(legacy), str(PRIV_OUT / w["slug"] / "week_manifest.json"))
    if (OUT / "common").exists():
        shutil.rmtree(OUT / "common")

    # 인덱스 단계
    meta = {w["slug"]: w for w in WEEKS}
    weeks = []
    for w in WEEKS:
        pth = PRIV_OUT / w["slug"] / "week_manifest.json"
        if not pth.exists():
            continue
        res = json.loads(pth.read_text(encoding="utf-8"))
        m = meta[w["slug"]]
        res["title"], res["topics"], res["reuse"], res["skip"], res["date"] = m["title"], m["topics"], m.get("reuse", []), m.get("skip", []), m["date"]
        weeks.append(res)
    CURRENT_WEEKS[:] = weeks
    partition_scope(weeks)
    common = process_common()
    compiled = compile_sections(weeks, args.max_tokens)
    write_stubs([e for w in weeks for e in w["entries"]] + compiled)
    for w in weeks:
        write_week_readme(w)
    write_index(weeks, common, compiled)
    write_llms_txt(weeks, compiled)
    cleanup_trees(weeks, common, compiled)
    rescrub_all()

    pub_weeks = [dict(w, entries=[e for e in w["entries"] if e.get("public")]) for w in weeks]
    MANIFEST.write_text(scrub_pii(json.dumps({"generated": TODAY, "max_tokens": args.max_tokens, "scope": "public",
                                              "weeks": pub_weeks, "compiled": compiled}, ensure_ascii=False, indent=1)), encoding="utf-8")
    (PRIV / "manifest.json").write_text(scrub_pii(json.dumps({"generated": TODAY, "max_tokens": args.max_tokens, "scope": "all",
                                                              "weeks": weeks, "common": common, "compiled": compiled}, ensure_ascii=False, indent=1)), encoding="utf-8")

    all_files = [f for w in weeks for e in w["entries"] for f in e["files"]] + [f for e in compiled for f in e["files"]]
    pub_files = [f for w in weeks for e in w["entries"] if e.get("public") for f in e["files"]] + [f for e in compiled for f in e["files"]]
    big = [f for f in all_files if f["tokens"] > args.max_tokens]
    print(f"\n완료: 주차 {len(weeks)}개, 문서 {len(all_files) + sum(len(e['files']) for e in common)}개 (공개 {len(pub_files)}개), 한도 초과 파일 {len(big)}개")
    for f in big:
        print("  초과:", f["path"], f["tokens"])


if __name__ == "__main__":
    main()
